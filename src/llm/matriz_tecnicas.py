"""Fase 30: matriz de técnicas (la tabla del profesor).

Filas = las 10 necesidades; columnas = Base | Few shot | RAG | Harness, por
modelo. Base y Few shot se miden; RAG y Harness quedan como «pendiente».

Cada celda (necesidad × técnica × modelo) corre R repeticiones de la cadena de
autocorrección de `iteraciones_juez` (árbitro: chequeo automático; corte al
primer acierto; temperatura 0.5) y guarda, con números reales:
  - termina_correcto: repeticiones cuya cadena terminó en «correcta» (x/R);
  - acierto_1er_intento: repeticiones que acertaron en el intento 1 (x/R): la
    métrica que mejor aísla el efecto del few-shot, porque no la contamina la
    autocorrección;
  - intento medio de convergencia (entre las que convergieron) y segundos
    medios por intento (pipeline + juez).

    python -m src.llm.matriz_tecnicas --modelos llama3.1:8b qwen2.5-coder:latest \\
        --tecnicas base few_shot --repeticiones 3 --iteraciones 3

Salidas: docs/resultados/matriz_tecnicas.json (versionado, lo lee el dashboard)
y docs/evidencia/fase30/matriz_tecnicas.{xlsx,csv,md}; los logs crudos de cada
repetición en docs/evidencia/fase30/logs/ (ignorado por git). El JSON se
reescribe después de cada repetición; con --reanudar se saltan las
repeticiones que ya estén en él.
"""
from __future__ import annotations

import argparse
import csv
import json
import statistics
import sys
from datetime import datetime
from pathlib import Path

from src.llm import iteraciones_juez as ij
from src.llm import tecnicas

RUTA_JSON = Path("docs/resultados/matriz_tecnicas.json")
DIR_EVIDENCIA = Path("docs/evidencia/fase30")
NECESIDADES = tuple(sorted(ij.PREGUNTAS, key=lambda n: int(n[1:])))
MODELOS_DEFAULT = ("llama3.1:8b", "qwen2.5-coder:latest")
ICONO = {ij.CORRECTA: "✓", ij.PARCIAL: "~", ij.INCORRECTA: "✗"}


def _secuencia(corrida: ij.CorridaIteraciones) -> str:
    return "".join(ICONO.get(it.veredicto, "?") for it in corrida.iteraciones)


def repeticion_de(corrida: ij.CorridaIteraciones) -> dict:
    """Lo que se guarda de una repetición (una cadena de autocorrección)."""
    its = corrida.iteraciones
    return {
        "secuencia": _secuencia(corrida),
        "termina_correcto": bool(its) and its[-1].veredicto == ij.CORRECTA,
        "acierto_1er_intento": bool(its) and its[0].veredicto == ij.CORRECTA,
        "convergio_en": corrida.convergencia.get(corrida.modelos[0]),
        "segundos_por_intento": [it.duracion_total_s for it in its],
        "ejemplos_intento_1": its[0].texto_ejemplos if its else "",
        "timestamp": corrida.timestamp,
    }


def celda(repeticiones: list[dict]) -> dict:
    """Métricas de una celda a partir de sus repeticiones."""
    r = len(repeticiones)
    conv = [x["convergio_en"] for x in repeticiones if x["convergio_en"]]
    segs = [s for x in repeticiones for s in x["segundos_por_intento"]]
    return {
        "repeticiones": r,
        "termina_correcto": sum(x["termina_correcto"] for x in repeticiones),
        "acierto_1er_intento": sum(x["acierto_1er_intento"] for x in repeticiones),
        "intento_medio_convergencia": round(statistics.mean(conv), 2) if conv else None,
        "seg_medio_intento": round(statistics.mean(segs), 1) if segs else None,
        "secuencias": [x["secuencia"] for x in repeticiones],
    }


def _clave(nid: str, modelo: str, tecnica: str) -> str:
    return f"{nid}|{modelo}|{tecnica}"


def _vacio(parametros: dict) -> dict:
    return {"generado": "", "parametros": parametros, "estado_tecnicas":
            {t: d.estado for t, d in tecnicas.TECNICAS.items()}, "repeticiones": {}}


def cargar(ruta: Path = RUTA_JSON) -> dict | None:
    try:
        return json.loads(ruta.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def celdas(data: dict) -> dict[str, dict]:
    """{clave «N1|modelo|tecnica»: métricas de la celda}."""
    return {k: celda(v) for k, v in data.get("repeticiones", {}).items()}


def totales(data: dict, modelo: str, tecnica: str) -> dict | None:
    """Suma de las celdas de un modelo y una técnica (fila TOTAL)."""
    cs = [c for k, c in celdas(data).items() if k.endswith(f"|{modelo}|{tecnica}")]
    if not cs:
        return None
    segs = [c["seg_medio_intento"] for c in cs if c["seg_medio_intento"] is not None]
    return {
        "repeticiones": sum(c["repeticiones"] for c in cs),
        "termina_correcto": sum(c["termina_correcto"] for c in cs),
        "acierto_1er_intento": sum(c["acierto_1er_intento"] for c in cs),
        "seg_medio_intento": round(statistics.mean(segs), 1) if segs else None,
    }


def texto_celda(c: dict | None, tecnica: str) -> str:
    if tecnicas.TECNICAS[tecnica].estado == tecnicas.PENDIENTE:
        return "pendiente"
    if not c:
        return "—"
    return f"termina {c['termina_correcto']}/{c['repeticiones']} · 1.er {c['acierto_1er_intento']}/{c['repeticiones']}"


def filas_matriz(data: dict, modelo: str) -> list[dict]:
    """Filas N1..N10 + TOTAL de un modelo, con una columna por técnica (las 4)."""
    cs = celdas(data)
    filas = []
    for nid in NECESIDADES:
        fila = {"Necesidad": nid}
        for t, d in tecnicas.TECNICAS.items():
            fila[d.etiqueta] = texto_celda(cs.get(_clave(nid, modelo, t)), t)
        filas.append(fila)
    fila = {"Necesidad": "TOTAL"}
    for t, d in tecnicas.TECNICAS.items():
        fila[d.etiqueta] = texto_celda(totales(data, modelo, t), t)
    filas.append(fila)
    return filas


def modelos_de(data: dict) -> list[str]:
    return list(dict.fromkeys(k.split("|")[1] for k in data.get("repeticiones", {})))


# ---------------------------------------------------------------------------
# Corrida
# ---------------------------------------------------------------------------

def guardar(data: dict, ruta: Path = RUTA_JSON) -> None:
    data["generado"] = datetime.now().astimezone().isoformat(timespec="seconds")
    data["celdas"] = celdas(data)
    data["totales"] = {f"{m}|{t}": totales(data, m, t) for m in modelos_de(data) for t in tecnicas.TECNICAS_ACTIVAS
                       if totales(data, m, t)}
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def correr(
    *, modelos: list[str], tecnicas_a_medir: list[str], necesidades: list[str], repeticiones: int, iteraciones: int,
    timeout_s: float, reanudar: bool, ruta: Path = RUTA_JSON, dir_logs: Path = DIR_EVIDENCIA / "logs",
) -> dict:
    for t in tecnicas_a_medir:
        tecnicas.validar_tecnica(t)
    parametros = {"modelos": modelos, "tecnicas": tecnicas_a_medir, "necesidades": necesidades,
                  "repeticiones": repeticiones, "iteraciones": iteraciones, "modo": ij.MODO_AUTOCORRECCION,
                  "arbitro": ij.ARBITRO_CHEQUEO, "corta_al_acertar": True, "medir_juez": False,
                  "temperatura": ij.TEMPERATURA_ITERACIONES, "timeout_llm_s": timeout_s}
    data = (cargar(ruta) if reanudar else None) or _vacio(parametros)
    data["parametros"] = parametros
    # Modelo por modelo: Ollama carga cada uno una sola vez.
    for modelo in modelos:
        for tecnica in tecnicas_a_medir:
            for nid in necesidades:
                hechas = data["repeticiones"].setdefault(_clave(nid, modelo, tecnica), [])
                for rep in range(len(hechas) + 1, repeticiones + 1):
                    corrida = ij.correr_iteraciones(
                        nid, modelos=[modelo], n=iteraciones, modo=ij.MODO_AUTOCORRECCION,
                        arbitro=ij.ARBITRO_CHEQUEO, medir_juez=False, corta_al_acertar=True, tecnica=tecnica,
                        timeout_sql_llm_s=timeout_s, timeout_respuesta_llm_s=timeout_s,
                    )
                    hechas.append(repeticion_de(corrida))
                    destino = dir_logs / modelo.replace(":", "_") / tecnica
                    destino.mkdir(parents=True, exist_ok=True)
                    (destino / f"{nid}_rep{rep}.json").write_text(ij.a_json(corrida), encoding="utf-8")
                    (destino / f"{nid}_rep{rep}.txt").write_text(ij.a_texto(corrida), encoding="utf-8")
                    guardar(data, ruta)
                    x = hechas[-1]
                    print(f"{modelo} · {tecnica} · {nid} · rep {rep}/{repeticiones}: {x['secuencia']} "
                          f"({sum(x['segundos_por_intento']):.0f} s) · ejemplos: {x['ejemplos_intento_1']}",
                          flush=True)
    guardar(data, ruta)
    return data


# ---------------------------------------------------------------------------
# Exportes (evidencia)
# ---------------------------------------------------------------------------

def filas_detalle(data: dict) -> list[dict]:
    salida = []
    for clave, reps in data.get("repeticiones", {}).items():
        nid, modelo, tecnica = clave.split("|")
        for i, x in enumerate(reps, 1):
            salida.append({"necesidad": nid, "modelo": modelo, "tecnica": tecnica, "repeticion": i,
                           "secuencia": x["secuencia"], "termina_correcto": x["termina_correcto"],
                           "acierto_1er_intento": x["acierto_1er_intento"], "convergio_en": x["convergio_en"],
                           "segundos_por_intento": " ".join(f"{s:.1f}" for s in x["segundos_por_intento"]),
                           "ejemplos_intento_1": x["ejemplos_intento_1"]})
    return salida


def exportar(data: dict, dir_salida: Path = DIR_EVIDENCIA) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    dir_salida.mkdir(parents=True, exist_ok=True)
    modelos = modelos_de(data)
    etiquetas = [d.etiqueta for d in tecnicas.TECNICAS.values()]
    # CSV y MD: una fila por necesidad, un bloque de columnas por modelo.
    encabezado = ["Necesidad"] + [f"{m} · {e}" for m in modelos for e in etiquetas]
    por_modelo = {m: filas_matriz(data, m) for m in modelos}
    filas = [[por_modelo[modelos[0]][i]["Necesidad"]] + [por_modelo[m][i][e] for m in modelos for e in etiquetas]
             for i in range(len(NECESIDADES) + 1)] if modelos else []
    ruta_csv = dir_salida / "matriz_tecnicas.csv"
    with ruta_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(encabezado)
        w.writerows(filas)
    ruta_md = dir_salida / "matriz_tecnicas.md"
    md = ["| " + " | ".join(encabezado) + " |", "|" + "---|" * len(encabezado)]
    md += ["| " + " | ".join(f) + " |" for f in filas]
    ruta_md.write_text(f"# Matriz de técnicas ({data.get('generado', '')})\n\n" + "\n".join(md) + "\n",
                       encoding="utf-8")
    # XLSX: hoja «Matriz» (encabezados naranja con texto blanco) y hoja «Detalle».
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Matriz"
    naranja = PatternFill("solid", fgColor="F28C28")
    blanco = Font(color="FFFFFF", bold=True)
    hoja.cell(row=1, column=1, value="Necesidad")
    col = 2
    for m in modelos:
        hoja.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(etiquetas) - 1)
        hoja.cell(row=1, column=col, value=m)
        for j, e in enumerate(etiquetas):
            hoja.cell(row=2, column=col + j, value=e)
        col += len(etiquetas)
    hoja.cell(row=2, column=1, value="")
    for fila in hoja.iter_rows(min_row=1, max_row=2, max_col=col - 1):
        for c in fila:
            c.fill, c.font, c.alignment = naranja, blanco, Alignment(horizontal="center")
    for i, f in enumerate(filas, start=3):
        for j, v in enumerate(f, start=1):
            hoja.cell(row=i, column=j, value=v)
    hoja.column_dimensions["A"].width = 11
    for j in range(2, col):
        hoja.column_dimensions[hoja.cell(row=2, column=j).column_letter].width = 24
    detalle = wb.create_sheet("Detalle")
    det = filas_detalle(data)
    if det:
        detalle.append(list(det[0]))
        for c in detalle[1]:
            c.fill, c.font = naranja, blanco
        for d in det:
            detalle.append(list(d.values()))
    ruta_xlsx = dir_salida / "matriz_tecnicas.xlsx"
    wb.save(ruta_xlsx)
    return [ruta_csv, ruta_md, ruta_xlsx]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m src.llm.matriz_tecnicas", description=__doc__.splitlines()[0])
    parser.add_argument("--modelos", nargs="+", default=list(MODELOS_DEFAULT))
    parser.add_argument("--tecnicas", nargs="+", default=list(tecnicas.TECNICAS_ACTIVAS),
                        choices=list(tecnicas.TECNICAS))
    parser.add_argument("--necesidades", nargs="+", default=list(NECESIDADES), choices=list(NECESIDADES))
    parser.add_argument("--repeticiones", type=int, default=3)
    parser.add_argument("--iteraciones", type=int, default=3)
    parser.add_argument("--timeout", type=float, default=300.0, help="Timeout (s) de cada llamada al modelo")
    parser.add_argument("--reanudar", action="store_true", help="Saltar las repeticiones ya guardadas en el JSON")
    parser.add_argument("--salida", default=str(RUTA_JSON), help="Ruta del JSON de la matriz")
    parser.add_argument("--solo-exportar", action="store_true", help="No correr nada: exportar el JSON existente")
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # ✓/✗ en la consola cp1252 de Windows
    ruta = Path(args.salida)
    if args.solo_exportar:
        data = cargar(ruta)
        if data is None:
            parser.error(f"No existe {ruta}")
    else:
        try:
            data = correr(modelos=args.modelos, tecnicas_a_medir=args.tecnicas, necesidades=args.necesidades,
                          repeticiones=args.repeticiones, iteraciones=args.iteraciones, timeout_s=args.timeout,
                          reanudar=args.reanudar, ruta=ruta)
        except tecnicas.TecnicaPendiente as exc:
            parser.error(str(exc))
    for p in exportar(data):
        print(f"Escrito: {p}")
    for m in modelos_de(data):
        print(f"\n{m}")
        for f in filas_matriz(data, m):
            print("  " + " | ".join(f"{k}: {v}" for k, v in f.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
