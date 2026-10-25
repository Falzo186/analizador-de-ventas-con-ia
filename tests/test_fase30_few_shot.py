"""Fase 30: few-shot con necesidades negativas, registro de técnicas,
trazabilidad, matriz de técnicas y validación fuera de muestra."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import duckdb
import pytest

from src.analytics.necesidades import DEFINICIONES
from src.llm import few_shot as fs
from src.llm import iteraciones_juez as ij
from src.llm import matriz_tecnicas as mt
from src.llm import responder_instacart as ri
from src.llm import retroalimentacion as rf
from src.llm import tecnicas
from src.llm import validacion_fuera_de_muestra as vfm
from src.llm.prompts_instacart import prompt_sql_instacart, prompt_sql_reintento_instacart
from src.llm.safety_instacart import validar_sql

DB = Path("data/instacart.duckdb")
REFERENCIA = Path("data/referencia")
hay_bd = pytest.mark.skipif(not DB.exists() or not REFERENCIA.exists(), reason="sin la BD real o sus referencias")


# ---------------------------------------------------------------------------
# Catálogo
# ---------------------------------------------------------------------------

def test_catalogo_cumple_el_esquema_y_tiene_los_14_ejemplos():
    ejs = fs.cargar_ejemplos()  # valida campos, tipos, ids únicos y verificación
    assert len(ejs) == 14
    assert len({e.id for e in ejs}) == 14
    por_necesidad = Counter(e.necesidad_origen for e in ejs)
    assert set(por_necesidad) == set(DEFINICIONES)
    for nid in DEFINICIONES:
        assert por_necesidad[nid] == (2 if nid in ("N2", "N4", "N5", "N10") else 1), nid
    for e in ejs:
        assert e.tipo == "negativa" and e.verificada and e.modelo_grande and e.fecha and e.resumen_resultado
        assert e.patron and e.palabras_clave


def test_cada_ejemplo_pasa_la_guarda_estatica():
    for e in fs.ejemplos():
        assert e.sql.lstrip().upper().startswith("SELECT"), e.id
        assert validar_sql(e.sql).startswith("SELECT"), e.id


def test_cargar_rechaza_ids_repetidos_y_campos_faltantes():
    crudos = json.loads(fs.RUTA_EJEMPLOS.read_text(encoding="utf-8"))
    with pytest.raises(ValueError, match="repetido"):
        fs.cargar_ejemplos(crudos + [crudos[0]])
    sin_sql = [dict(crudos[0], sql=None)]
    with pytest.raises(ValueError, match="sql"):
        fs.cargar_ejemplos(sin_sql)


def test_ningun_ejemplo_usa_como_ancla_la_respuesta_de_una_necesidad():
    anclas = ("46408", "dairy eggs", "Raw Veggie Wrappers", "Organic Hass Avocado", "Bag of Organic Bananas",
              "21.49", "11.87", "288418", "64761")
    for e in fs.ejemplos():
        texto = f"{e.pregunta} {e.sql} {e.patron} {' '.join(e.palabras_clave)}"
        for a in anclas:
            assert a.lower() not in texto.lower(), (e.id, a)


@hay_bd
def test_cada_ejemplo_ejecuta_y_su_resultado_difiere_de_la_referencia():
    from scripts.verificar_few_shot import verificar
    for e in json.loads(fs.RUTA_EJEMPLOS.read_text(encoding="utf-8")):
        r = verificar(e)
        assert r["ok"], (e["id"], r["motivo"])


# ---------------------------------------------------------------------------
# Recuperación y anti-fuga del bloque inyectado
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("nid", sorted(DEFINICIONES, key=lambda n: int(n[1:])))
def test_top1_de_cada_necesidad_es_su_hermana(nid):
    recuperados = fs.recuperar(DEFINICIONES[nid].pregunta)
    assert recuperados, nid
    assert recuperados[0][0].necesidad_origen == nid
    assert all(s >= fs.UMBRAL_SIMILITUD for _, s in recuperados)
    assert len(recuperados) <= fs.K_DEFAULT


def test_pregunta_ajena_no_recibe_ejemplos():
    assert fs.recuperar("¿cuál es la capital de Francia?") == []
    ctx = fs.contexto_few_shot("¿cuál es la capital de Francia?")
    assert ctx.bloque is None and ctx.usados == []


def test_k_por_variable_de_entorno(monkeypatch):
    monkeypatch.setenv("VENTIA_FEWSHOT_K", "1")
    assert len(fs.recuperar(DEFINICIONES["N4"].pregunta)) == 1
    monkeypatch.setenv("VENTIA_FEWSHOT_K", "3")
    assert len(fs.recuperar(DEFINICIONES["N4"].pregunta)) == 3


def test_bloque_lleva_pregunta_sql_y_patron_pero_no_el_resultado():
    ctx = fs.contexto_few_shot(DEFINICIONES["N5"].pregunta)
    e = fs.recuperar(DEFINICIONES["N5"].pregunta)[0][0]
    assert ctx.bloque.startswith(fs.ENCABEZADO_BLOQUE)
    assert e.pregunta in ctx.bloque and e.sql in ctx.bloque and e.patron in ctx.bloque
    for fila in e.resumen_resultado:
        for v in fila.values():
            assert str(v) not in ctx.bloque, v
    assert ctx.usados[0][0] == e.id


@hay_bd
@pytest.mark.parametrize("nid", sorted(DEFINICIONES, key=lambda n: int(n[1:])))
def test_el_bloque_no_contiene_el_valor_clave_de_su_necesidad(nid):
    # Mismo enfoque que el test de no-fuga de la Fase 24 (valores_protegidos + fugas).
    ref = json.loads((REFERENCIA / f"{nid}.json").read_text(encoding="utf-8"))
    pregunta = DEFINICIONES[nid].pregunta
    vp = rf.valores_protegidos(pregunta, [{f"v{i}": vc["valor"] for i, vc in enumerate(ref["valores_clave"])}])
    bloque = fs.contexto_few_shot(pregunta).bloque
    assert rf.fugas(bloque, vp) == [], nid


# ---------------------------------------------------------------------------
# Técnicas y prompt
# ---------------------------------------------------------------------------

GOLDEN = json.loads(Path("tests/data/prompt_base_golden.json").read_text(encoding="utf-8"))
P = "¿Qué departamento tiene la tasa de recompra promedio más alta?"
E = "ESQUEMA DE PRUEBA\n- orders: order_id INTEGER"


def test_tecnica_base_deja_el_prompt_identico_byte_por_byte():
    assert list(prompt_sql_instacart(P, E)) == GOLDEN["sql"]
    assert list(prompt_sql_instacart(P, E, None, "RETRO", "LECCIONES", None)) == GOLDEN["sql_lecciones_retro"]
    assert list(prompt_sql_reintento_instacart(P, E, "SELECT 1", "MOTIVO")) == GOLDEN["reintento"]
    assert list(prompt_sql_reintento_instacart(P, E, "SELECT 1", "MOTIVO", None, "RETRO", "LECCIONES", None)) \
        == GOLDEN["reintento_lecciones"]
    assert tecnicas.contexto_de_tecnica(tecnicas.BASE, P).bloque is None


def test_bloque_va_despues_del_esquema_y_antes_de_la_bitacora_y_la_pregunta():
    bloque = fs.contexto_few_shot(P).bloque
    sistema, usuario = prompt_sql_instacart(P, E, None, None, "LECCIONES DE LA BITACORA", bloque)
    assert sistema.index(E) < sistema.index(fs.ENCABEZADO_BLOQUE) < sistema.index("LECCIONES DE LA BITACORA")
    assert fs.ENCABEZADO_BLOQUE not in usuario and P in usuario


def test_rag_y_harness_estan_registradas_como_pendientes():
    assert tecnicas.TECNICAS_ACTIVAS == (tecnicas.BASE, tecnicas.FEW_SHOT)
    for t in (tecnicas.RAG, tecnicas.HARNESS):
        assert tecnicas.TECNICAS[t].estado == tecnicas.PENDIENTE
        with pytest.raises(tecnicas.TecnicaPendiente, match="pendiente"):
            tecnicas.validar_tecnica(t)
    with pytest.raises(ValueError, match="desconocida"):
        tecnicas.validar_tecnica("magia")
    with pytest.raises(tecnicas.TecnicaPendiente):
        ij.nueva_corrida("N5", modelos=["m"], tecnica=tecnicas.RAG)


def test_default_es_base_y_se_cambia_por_entorno(monkeypatch):
    monkeypatch.delenv("VENTIA_TECNICA", raising=False)
    assert tecnicas.TECNICA_DEFAULT == tecnicas.BASE == tecnicas.tecnica_por_defecto()
    monkeypatch.setenv("VENTIA_TECNICA", "few_shot")
    assert tecnicas.tecnica_por_defecto() == tecnicas.FEW_SHOT


def _pipeline_sin_modelo(monkeypatch, sqls_vistos: list):
    monkeypatch.setattr(ri, "abrir_conexion_solo_lectura", lambda _: duckdb.connect())
    monkeypatch.setattr(ri, "construir_descripcion_esquema", lambda con: E)

    def generar(pregunta, esquema, **kwargs):
        sqls_vistos.append(kwargs.get("ejemplos"))
        return "SELECT 1 AS n"
    monkeypatch.setattr(ri, "generar_sql_instacart", generar)


def test_responder_inyecta_los_ejemplos_solo_con_few_shot(monkeypatch):
    vistos: list = []
    _pipeline_sin_modelo(monkeypatch, vistos)
    base = ri.responder_pregunta_instacart(DEFINICIONES["N5"].pregunta, redactar=False, tecnica="base")
    few = ri.responder_pregunta_instacart(DEFINICIONES["N5"].pregunta, redactar=False, tecnica="few_shot")
    assert base.ok and few.ok
    assert vistos[0] is None and base.tecnica == "base" and base.ejemplos_usados == []
    assert vistos[1] == few.bloque_ejemplos and few.tecnica == "few_shot"
    assert few.ejemplos_usados[0][0] == "N5-neg-1"
    pendiente = ri.responder_pregunta_instacart("p", redactar=False, tecnica="rag")
    assert not pendiente.ok and "pendiente" in pendiente.error


# ---------------------------------------------------------------------------
# Trazabilidad y métricas por intento
# ---------------------------------------------------------------------------

def _fake_iteraciones(monkeypatch, veredictos: list[bool]):
    llamadas = iter(veredictos)

    def pipeline(pregunta, **kwargs):
        res = ri.ResultadoInstacart(pregunta=pregunta, sql="SELECT 1", columnas=["x"], filas=[(1,)],
                                    respuesta="r", tecnica=kwargs["tecnica"])
        if kwargs["tecnica"] == "few_shot":
            res.ejemplos_usados, res.bloque_ejemplos = [("N5-neg-1", 0.62)], "BLOQUE"
        res.filas = [(10.1,)] if next(llamadas) else [(99.0,)]
        return res

    monkeypatch.setattr(ri, "responder_pregunta_instacart", pipeline)
    monkeypatch.setattr(ij, "diagnosticar",
                        lambda *a, **k: ij.Diagnostico("comparación", "agrupa por pedido", "juez", "{}"))


def test_cada_intento_registra_tecnica_ejemplos_y_si_mejoro(monkeypatch):
    _fake_iteraciones(monkeypatch, [False, True])
    corrida = ij.CorridaIteraciones(
        necesidad="N5", pregunta=DEFINICIONES["N5"].pregunta, referencia="ref", modelos=["m"], n_iteraciones=2,
        modo=ij.MODO_AUTOCORRECCION, arbitro=ij.ARBITRO_CHEQUEO, medir_juez=False, corta_al_acertar=False,
        tecnica="few_shot", referencia_filas=[{"promedio_productos_por_pedido": 10.1}],
    )
    ij.correr_iteraciones("N5", modelos=["m"], corrida=corrida)
    primero, segundo = corrida.iteraciones
    assert (primero.tecnica, primero.ejemplos_usados, primero.bloque_ejemplos) == ("few_shot", [["N5-neg-1", 0.62]],
                                                                                   "BLOQUE")
    assert [f["cambio"] for f in corrida.tabla()] == [None, ij.MEJORO]
    assert corrida.tabla()[0]["pista_del_juez"] == "agrupa por pedido"
    texto = ij.a_texto(corrida)
    assert "Técnica de contexto: few_shot" in texto
    assert "Técnica: few_shot — ejemplos usados: N5-neg-1 (sim 0.62)" in texto
    assert "Cambio respecto al intento anterior: mejoró" in texto
    d = corrida.a_dict()
    assert d["tecnica"] == "few_shot" and d["iteraciones"][0]["ejemplos_usados"] == [["N5-neg-1", 0.62]]
    assert d["tabla_iteraciones"][1]["cambio"] == ij.MEJORO
    r = corrida.resumen()["m"]
    assert (r.pct_correctas, r.pct_acercamiento) == (50.0, 50.0)


def test_tendencia_y_porcentaje_de_acercamiento():
    assert ij.tendencia(None, ij.CORRECTA) is None
    assert ij.tendencia(ij.INCORRECTA, ij.PARCIAL) == ij.MEJORO
    assert ij.tendencia(ij.CORRECTA, ij.PARCIAL) == ij.EMPEORO
    assert ij.tendencia(ij.PARCIAL, ij.PARCIAL) == ij.IGUAL
    r = ij.ResumenModelo("m", 4, 1, 2, 1, 0)
    assert (r.pct_correctas, r.pct_acercamiento) == (25.0, 75.0)


# ---------------------------------------------------------------------------
# Matriz de técnicas
# ---------------------------------------------------------------------------

def _datos_matriz() -> dict:
    rep = {"secuencia": "✗✓", "termina_correcto": True, "acierto_1er_intento": False, "convergio_en": 2,
           "segundos_por_intento": [4.0, 6.0], "ejemplos_intento_1": "ninguno", "timestamp": ""}
    rep1 = dict(rep, secuencia="✓", acierto_1er_intento=True, convergio_en=1, segundos_por_intento=[2.0])
    return {"parametros": {}, "repeticiones": {"N5|m|base": [rep], "N5|m|few_shot": [rep1, rep1]}}


def test_matriz_celdas_totales_y_pendientes():
    data = _datos_matriz()
    c = mt.celdas(data)
    assert c["N5|m|base"]["termina_correcto"] == 1 and c["N5|m|base"]["acierto_1er_intento"] == 0
    assert c["N5|m|base"]["intento_medio_convergencia"] == 2 and c["N5|m|base"]["seg_medio_intento"] == 5.0
    assert mt.totales(data, "m", "few_shot")["acierto_1er_intento"] == 2
    filas = mt.filas_matriz(data, "m")
    assert [f["Necesidad"] for f in filas] == [*mt.NECESIDADES, "TOTAL"]
    n5 = filas[4]
    assert n5["Base"] == "termina 1/1 · 1.er 0/1" and n5["Few shot"] == "termina 2/2 · 1.er 2/2"
    assert n5["RAG"] == n5["Harness"] == "pendiente"
    assert filas[0]["Base"] == "—"


def test_matriz_exporta_csv_md_y_xlsx(tmp_path):
    from openpyxl import load_workbook
    rutas = mt.exportar(_datos_matriz(), tmp_path)
    assert {p.suffix for p in rutas} == {".csv", ".md", ".xlsx"}
    wb = load_workbook(tmp_path / "matriz_tecnicas.xlsx")
    assert wb.sheetnames == ["Matriz", "Detalle"]
    hoja = wb["Matriz"]
    assert hoja.cell(row=1, column=2).value == "m" and hoja.cell(row=2, column=4).value == "RAG"
    assert hoja.cell(row=1, column=1).fill.fgColor.rgb.endswith("F28C28")
    assert hoja.cell(row=1, column=1).font.color.rgb.endswith("FFFFFF")
    assert hoja.cell(row=13, column=1).value == "TOTAL"


# ---------------------------------------------------------------------------
# Validación fuera de muestra (adaptador del comparador)
# ---------------------------------------------------------------------------

def test_adaptador_del_comparador_para_preguntas_libres():
    preguntas = {p["id"]: p for p in vfm.cargar_preguntas()}
    assert len(preguntas) == 8 and "¿Qué departamento tiene más pedidos?" == preguntas["X1"]["pregunta"]
    ref_x1 = [{"department": "produce", "pedidos": 2506247}, {"department": "dairy eggs", "pedidos": 2264738}]
    forma = vfm.forma_de(preguntas["X1"])
    assert rf.coincide_con_forma(forma, ref_x1, [("produce", 2506247)])[0]
    assert not rf.coincide_con_forma(forma, ref_x1, [("dairy eggs", 2264738)])[0]
    ref_x2 = [{"usuarios": 206209}]
    assert rf.coincide_con_forma(vfm.forma_de(preguntas["X2"]), ref_x2, [(206209,)])[0]
    assert not rf.coincide_con_forma(vfm.forma_de(preguntas["X2"]), ref_x2, [(3421083,)])[0]


def test_resumen_fuera_de_muestra_detecta_empeoradas():
    filas = [{"id": "X1", "base": {"acierta": True}, "few_shot": {"acierta": False, "ejemplos_usados": []}},
             {"id": "X2", "base": {"acierta": False}, "few_shot": {"acierta": True, "ejemplos_usados": [["a", 0.5]]}}]
    r = vfm.resumir(filas)
    assert r["empeoradas_por_few_shot"] == ["X1"] and r["mejoradas_por_few_shot"] == ["X2"]
    assert r["reciben_ejemplos"] == ["X2"]


@hay_bd
def test_los_sql_de_referencia_fuera_de_muestra_dan_lo_esperado():
    p = {x["id"]: x for x in vfm.cargar_preguntas()}
    assert vfm.filas_referencia(p["X1"])[0] == {"department": "produce", "pedidos": 2506247}
    assert vfm.filas_referencia(p["X2"])[0] == {"usuarios": 206209}


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

from tests.test_app_instacart import ModeloSimulado, app, rutas  # noqa: E402,F401 - fixtures de la app


def test_ui_selector_de_tecnica_y_secciones_nuevas(app):  # noqa: F811
    at = app(ModeloSimulado([], []))
    selector = at.selectbox(key="tecnica_sel")
    assert list(selector.options) == ["Base", "Few shot"] and selector.value == "base"
    assert not at.radio, "en la vista de chat no debe haber radios (el radio es el selector del dashboard)"
    pendientes = [c for c in at.checkbox if (c.key or "").startswith("tecnica_pendiente_")]
    assert len(pendientes) == 2 and all(c.disabled for c in pendientes)
    assert all("pendiente (siguiente fase)" in c.label for c in pendientes)
    at.button(key="vista_dashboard").click().run()
    assert not list(at.exception)
    etiquetas = [e.label for e in at.expander]
    assert any("Necesidades negativas" in x for x in etiquetas)
    assert any("Matriz de técnicas" in x for x in etiquetas)
